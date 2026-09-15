from datetime import datetime
from io import BytesIO

import os
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from fastapi import APIRouter, BackgroundTasks, Depends, Query, HTTPException, Response
from fastapi.responses import StreamingResponse
from sqlalchemy import extract
from sqlalchemy.orm import Session
from typing import Optional
from pydantic import BaseModel, Field

from ..database import get_db
from ..models import ConferenceEvent
from ..schemas import ConferenceEventOut
from .reports import DAY_FULL, _weekday_idx

_FRONTEND_URL = os.environ.get("FRONTEND_URL", "http://localhost:4000")

# ต้องตรงกับ label ใน Frontend/src/components/EventForm.jsx (chips ของฟิลด์ status)
_STATUS_READY = "สร้าง link แล้ว"

router = APIRouter(prefix="/events", tags=["Events"])


def _escape_like(value: str) -> str:
    """Escape LIKE/ILIKE wildcards (\\, %, _) ในค่าที่ผู้ใช้ส่งมา ก่อนนำไปสร้าง pattern"""
    return value.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")


def get_event_or_404(event_id: int, db: Session = Depends(get_db)) -> ConferenceEvent:
    ev = db.query(ConferenceEvent).filter(ConferenceEvent.id == event_id).first()
    if not ev:
        raise HTTPException(status_code=404, detail="ไม่พบรายการนี้")
    return ev


# ---------------------------------------------------------------------------
# Schema สำหรับ create / update
# ---------------------------------------------------------------------------
class EventIn(BaseModel):
    # ความยาวสูงสุดตรงกับ column width ใน models.py — กัน request ที่เกิน limit
    # ให้ตอบ 422 ที่ชัดเจน แทนที่จะไปพังตอน insert ลง MySQL
    date:         Optional[str] = Field(default=None, max_length=10)
    time_raw:     Optional[str] = Field(default=None, max_length=255)
    title:        Optional[str] = Field(default=None, max_length=10000)
    location:     Optional[str] = Field(default=None, max_length=255)
    app:          Optional[str] = Field(default=None, max_length=255)
    coordinator:  Optional[str] = Field(default=None, max_length=255)
    department:   Optional[str] = Field(default=None, max_length=255)
    book_no:      Optional[str] = Field(default=None, max_length=255)
    status:       Optional[str] = Field(default=None, max_length=100)
    assignee:     Optional[str] = Field(default=None, max_length=255)
    zoom_user:    Optional[str] = Field(default=None, max_length=255)
    details:      Optional[str] = Field(default=None, max_length=10000)
    month_source: Optional[str] = Field(default=None, max_length=50)
    image_url:    Optional[str] = Field(default=None, max_length=500)
    meeting_link: Optional[str] = Field(default=None, max_length=500)
    event_type:   Optional[str] = Field(default=None, max_length=50)

# ค่า default ตอนสร้างรายการใหม่ — ใช้แทนฟิลด์ที่ frontend ไม่ได้ส่งมา
_CREATE_DEFAULTS = {
    "date": "", "time_raw": "", "title": "", "location": "",
    "app": "Zoom", "coordinator": "", "department": "", "book_no": "",
    "status": "", "assignee": "", "zoom_user": "", "details": "", "month_source": "", "image_url": "",
    "meeting_link": "", "event_type": "",
}


# ---------------------------------------------------------------------------
# GET /events — ดึงทั้งหมด (ไม่รวมที่ถูกลบ)
# ---------------------------------------------------------------------------
_EVENTS_LIMIT = 5000


@router.get("/", response_model=list[ConferenceEventOut])
def get_events(
    response: Response,
    month:    Optional[str] = Query(None),
    date:     Optional[str] = Query(None),
    app:      Optional[str] = Query(None),
    assignee: Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    query = db.query(ConferenceEvent).filter(ConferenceEvent.deleted_at.is_(None))
    if month:    query = query.filter(ConferenceEvent.month_source == month)
    if date:     query = query.filter(ConferenceEvent.date == date)
    if app:      query = query.filter(ConferenceEvent.app.ilike(f"%{_escape_like(app)}%", escape="\\"))
    if assignee: query = query.filter(ConferenceEvent.assignee.ilike(f"%{_escape_like(assignee)}%", escape="\\"))
    rows = query.order_by(ConferenceEvent.date).limit(_EVENTS_LIMIT + 1).all()
    # ถ้าจำนวนแถวเกิน limit จริง ให้ตัดกลับเหลือ limit แล้วแจ้งผ่าน header ว่าผลลัพธ์ถูกตัด
    # เพื่อไม่ให้ frontend เข้าใจผิดว่าได้ข้อมูลครบ (แทนที่จะเงียบๆ ตัดทิ้งแบบเดิม)
    if len(rows) > _EVENTS_LIMIT:
        rows = rows[:_EVENTS_LIMIT]
        response.headers["X-Results-Truncated"] = "true"
    return rows


# ---------------------------------------------------------------------------
# GET /events/export-excel — ส่งออกเป็นไฟล์ Excel
# ---------------------------------------------------------------------------
_EXPORT_HEADERS = [
    "วันที่", "วัน", "เวลา", "ชื่อการประชุม", "หน่วยงาน", "สถานที่", "App",
    "ผู้ประสานงาน", "ผู้รับผิดชอบ", "สถานะ", "บัญชี Zoom/Teams", "เลขที่หนังสือ", "รายละเอียด",
]
_EXPORT_COL_WIDTHS = [14, 14, 16, 40, 24, 20, 10, 20, 18, 16, 20, 16, 30]


@router.get("/export-excel")
def export_excel(
    year: Optional[int] = Query(None),
    month: Optional[int] = Query(None),
    week_start: Optional[str] = Query(None),
    week_end: Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    q = db.query(ConferenceEvent).filter(ConferenceEvent.deleted_at.is_(None))

    if week_start and week_end:
        q = q.filter(ConferenceEvent.date >= week_start).filter(ConferenceEvent.date <= week_end)
    elif year:
        q = q.filter(extract("year", ConferenceEvent.date) == year)
        if month:
            q = q.filter(extract("month", ConferenceEvent.date) == month)

    events = q.order_by(ConferenceEvent.date).limit(5000).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ตารางประชุม"

    header_fill = PatternFill("solid", fgColor="1B5E20")
    header_font = Font(bold=True, color="FFFFFF", name="Sarabun", size=11)
    stripe_fill = PatternFill("solid", fgColor="E8F5E9")

    for col, h in enumerate(_EXPORT_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row, ev in enumerate(events, 2):
        day_name = ""
        if ev.date:
            idx = _weekday_idx(ev.date)
            if idx is not None:
                day_name = DAY_FULL[idx]

        ws.append([
            str(ev.date) if ev.date else "",
            day_name,
            ev.time_raw or "",
            ev.title or "",
            ev.department or "",
            ev.location or "",
            ev.app or "",
            ev.coordinator or "",
            ev.assignee or "",
            ev.status or "",
            ev.zoom_user or "",
            ev.book_no or "",
            ev.details or "",
        ])
        if row % 2 == 0:
            for col in range(1, len(_EXPORT_HEADERS) + 1):
                ws.cell(row=row, column=col).fill = stripe_fill

    for col, width in enumerate(_EXPORT_COL_WIDTHS, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=conference_schedule.xlsx"},
    )


# ---------------------------------------------------------------------------
# ⚠️ Route เฉพาะทาง (all/, trash/) ต้องอยู่ก่อน /{event_id} เสมอ
# ไม่งั้น FastAPI จะจับ "all" หรือ "trash" เป็นค่า event_id (int) แล้ว error
# ---------------------------------------------------------------------------

@router.delete("/all/soft-delete-all")
def soft_delete_all(confirm: str = Query(...), db: Session = Depends(get_db)):
    if confirm != "DELETE_ALL_CONFIRMED":
        raise HTTPException(status_code=400, detail="ต้องยืนยันด้วย confirm=DELETE_ALL_CONFIRMED")

    count = (
        db.query(ConferenceEvent)
        .filter(ConferenceEvent.deleted_at.is_(None))
        .update({ConferenceEvent.deleted_at: datetime.utcnow()})
    )
    db.commit()
    return {"success": True, "message": f"ลบข้อมูลทั้งหมด {count} รายการแล้ว (กู้คืนได้ในถังขยะ)", "count": count}


@router.post("/all/restore-all")
def restore_all(db: Session = Depends(get_db)):
    count = (
        db.query(ConferenceEvent)
        .filter(ConferenceEvent.deleted_at.isnot(None))
        .update({ConferenceEvent.deleted_at: None})
    )
    db.commit()
    return {"success": True, "message": f"กู้คืนข้อมูลทั้งหมด {count} รายการแล้ว", "count": count}


@router.get("/trash/list", response_model=list[ConferenceEventOut])
def get_trash(db: Session = Depends(get_db)):
    return (
        db.query(ConferenceEvent)
        .filter(ConferenceEvent.deleted_at.isnot(None))
        .order_by(ConferenceEvent.deleted_at.desc())
        .limit(1000)
        .all()
    )


# ---------------------------------------------------------------------------
# GET /events/{id} — ดูรายการเดียว
# ---------------------------------------------------------------------------
@router.get("/{event_id}", response_model=ConferenceEventOut)
def get_event(ev: ConferenceEvent = Depends(get_event_or_404)):
    return ev


# ---------------------------------------------------------------------------
# POST /events — เพิ่มรายการใหม่
# ---------------------------------------------------------------------------
@router.post("/", response_model=ConferenceEventOut)
def create_event(body: EventIn, background_tasks: BackgroundTasks, db: Session = Depends(get_db)):
    if not body.title or not body.date:
        raise HTTPException(status_code=400, detail="กรุณากรอกชื่อเรื่องและวันที่")
    data = {k: (v if v is not None else _CREATE_DEFAULTS[k]) for k,v in body.model_dump().items()}
    ev = ConferenceEvent(**data)
    db.add(ev)
    db.commit()
    db.refresh(ev)

    from ..services.line_service import send_line_flex_new
    background_tasks.add_task(send_line_flex_new, ev)

    return ev


# ---------------------------------------------------------------------------
# PUT /events/{id} — แก้ไขรายการ
# ---------------------------------------------------------------------------
@router.put("/{event_id}", response_model=ConferenceEventOut)
def update_event(event_id: int, body: EventIn, background_tasks: BackgroundTasks, db: Session = Depends(get_db)):
    ev = db.query(ConferenceEvent).filter(ConferenceEvent.id == event_id).first()
    if not ev:
        raise HTTPException(status_code=404, detail="ไม่พบรายการนี้")

    _WATCH_FIELDS = ["title","date","time_raw","location","department",
                     "coordinator","assignee","app","meeting_link","book_no","status","details"]
    old_values = {k: getattr(ev, k) for k in _WATCH_FIELDS}

    for key, val in body.model_dump(exclude_unset=True).items():
        setattr(ev, key, "" if val is None else val)
    db.commit()
    db.refresh(ev)

    #แจ้งเตือนเมื่อมีฟิลด์ที่ติดตามเปลี่ยนแปลงจริง
    changes = {
        k: (old_values[k], getattr(ev, k))
        for k in _WATCH_FIELDS
        if str(old_values[k] or "") != str(getattr(ev, k) or "")
    }
    if changes:
        from ..services.line_service import send_line_flex_change
        background_tasks.add_task(send_line_flex_change, ev, changes)

    # แจ้งเตือนตอน status เปลี่ยน "เข้าสู่" สถานะพร้อมแล้ว (ไม่ยิงซ้ำถ้า save ซ้ำตอนสถานะนี้อยู่แล้ว)
    became_ready = old_values["status"] != _STATUS_READY and ev.status == _STATUS_READY
    if became_ready:
        from ..services.line_service import send_line_flex_ready
        background_tasks.add_task(send_line_flex_ready, ev)

    return ev


# ---------------------------------------------------------------------------
# DELETE /events/{id} — ลบรายการ (soft delete)
# ---------------------------------------------------------------------------
@router.delete("/{event_id}")
def delete_event(event_id: int, background_tasks: BackgroundTasks, db: Session = Depends(get_db)):
    ev = db.query(ConferenceEvent).filter(ConferenceEvent.id == event_id).first()
    if not ev:
        raise HTTPException(status_code=404, detail="ไม่พบรายการนี้")

    ev.deleted_at = datetime.utcnow()
    db.commit()

    from ..services.line_service import send_line_flex_delete
    background_tasks.add_task(send_line_flex_delete, ev)

    return {"success": True, "message": "ลบรายการสำเร็จ (กู้คืนได้ในถังขยะ)"}





# ---------------------------------------------------------------------------
# POST /events/{id}/restore — กู้คืนรายการเดียว
# ---------------------------------------------------------------------------
@router.post("/{event_id}/restore", response_model=ConferenceEventOut)
def restore_event(db: Session = Depends(get_db), ev: ConferenceEvent = Depends(get_event_or_404)):
    ev.deleted_at = None
    db.commit()
    db.refresh(ev)
    return ev


# ---------------------------------------------------------------------------
# DELETE /events/{id}/permanent — ลบถาวร (ไม่สามารถกู้คืนได้)
# ---------------------------------------------------------------------------
@router.delete("/{event_id}/permanent")
def permanent_delete_event(db: Session = Depends(get_db), ev: ConferenceEvent = Depends(get_event_or_404)):
    db.delete(ev)
    db.commit()
    return {"success": True, "message": "ลบถาวรแล้ว ไม่สามารถกู้คืนได้"}